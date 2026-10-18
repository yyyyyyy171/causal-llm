"""
周报生成器。

按 JSON 规格文件生成 .docx，风格与 E:\\汇报 里已有的周报一致
（标题 + 分节 + 表格）。规格文件里可以写固定文字，也可以直接引用
results/ 下的 CSV，让表格数据永远来自真实实验结果，而不是手抄。

用法：
    python scripts/make_week_report.py --spec scripts/week_reports/week4.json

规格文件支持的节类型：
    {"type": "heading",  "text": "..."}
    {"type": "para",     "text": "..."}
    {"type": "bullets",  "items": ["...", "..."]}
    {"type": "table",    "header": [...], "rows": [[...], ...]}
    {"type": "csv_table","path": "results/xxx.csv", "columns": [...],
                         "filter": {"col": "value"}, "sort_by": "col"}
    {"type": "image",    "path": "results/xxx.png", "width_inches": 6.0}
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import docx
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def resolve(path: str) -> Path:
    candidate = Path(path)
    if candidate.is_absolute():
        return candidate
    return PROJECT_ROOT / candidate


def read_csv_table(spec: dict):
    path = resolve(spec["path"])
    with open(path, "r", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))

    for column, value in (spec.get("filter") or {}).items():
        rows = [r for r in rows if str(r.get(column)) == str(value)]

    if spec.get("sort_by"):
        keys = spec["sort_by"]
        if isinstance(keys, str):
            keys = [keys]
        rows.sort(key=lambda r: tuple(_sort_key(r[k]) for k in keys))

    if spec.get("round"):
        digits = int(spec["round"])
        for row in rows:
            for key, value in row.items():
                try:
                    number = float(value)
                except (TypeError, ValueError):
                    continue
                # 只在“原文本来就是整数写法”时保持整数（样本量、计数、边数），
                # 免得出现 100.0000。
                # 原文写的是 0.0 / 1.0 时必须按位数格式化：否则会被压成 0 和 1，
                # 在一张 0 到 1 的指标表里看起来像另一回事（曾经就是这样，
                # 1.000 的准确率显示成 1，0.000 显示成 0）。
                raw = str(value)
                if ("." not in raw) and ("e" not in raw.lower()):
                    row[key] = str(int(number))
                else:
                    row[key] = f"{number:.{digits}f}"

    columns = spec.get(
        "columns", list(rows[0].keys()) if rows else []
    )

    return list(columns), [[row.get(c, "") for c in columns] for row in rows]


def _sort_key(value):
    try:
        return (0, float(value))
    except (TypeError, ValueError):
        return (1, str(value))


def add_table(document, header, rows):
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"

    for cell, text in zip(table.rows[0].cells, header):
        cell.text = str(text)
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)

    return table


def build(spec: dict, out_path: Path) -> None:

    document = docx.Document()

    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)

    title = spec.get("title")
    if title:
        heading = document.add_heading(title, level=1)
        for run in heading.runs:
            run.font.name = "Microsoft YaHei"

    subtitle = spec.get("subtitle")
    if subtitle:
        document.add_paragraph(subtitle)

    for section in spec.get("sections", []):

        kind = section["type"]

        if kind == "heading":
            document.add_heading(section["text"], level=2)

        elif kind == "para":
            document.add_paragraph(section["text"])

        elif kind == "bullets":
            for item in section["items"]:
                document.add_paragraph(item, style="List Bullet")

        elif kind == "table":
            add_table(document, section["header"], section["rows"])

        elif kind == "csv_table":
            header, rows = read_csv_table(section)
            if not rows:
                document.add_paragraph(f"（{section['path']} 暂无数据）")
                continue
            add_table(document, header, rows)

        elif kind == "image":
            path = resolve(section["path"])
            if path.exists():
                document.add_picture(
                    str(path),
                    width=Inches(float(section.get("width_inches", 6.0))),
                )
            else:
                document.add_paragraph(f"（缺少图片 {section['path']}）")

        else:
            raise ValueError(f"未知的节类型: {kind}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))

    print(f"周报已生成: {out_path}")


def main() -> int:

    parser = argparse.ArgumentParser()
    parser.add_argument("--spec", required=True)
    parser.add_argument("--out", default="")

    args = parser.parse_args()

    spec = json.loads(resolve(args.spec).read_text(encoding="utf-8"))

    out = args.out or spec.get("out")
    if not out:
        raise SystemExit("请在规格文件里写 out，或用 --out 指定输出路径")

    build(spec, Path(out))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
